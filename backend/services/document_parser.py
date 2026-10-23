import os
import shutil
from typing import Dict, List, Tuple

from docx import Document
from pypdf import PdfReader


MAX_TEXT_LENGTH = 2_000_000
MAX_PDF_PAGES = 60
MIN_PAGE_TEXT_LENGTH = 40
OCR_DPI = 180
OCR_PAGE_TIMEOUT_SECONDS = 30
PDF_RENDER_TIMEOUT_SECONDS = 35


class DocumentExtractionError(ValueError):
    """An input or environment problem that can be reported to the uploader."""


def _validate_file_path(file_path: str) -> str:
    normalized = os.path.abspath(file_path)
    if not os.path.exists(normalized):
        raise DocumentExtractionError("The uploaded document could not be found.")
    if not os.path.isfile(normalized):
        raise DocumentExtractionError("The uploaded document path is invalid.")
    if os.path.basename(normalized) != os.path.basename(file_path):
        raise DocumentExtractionError("Unsafe file name detected.")
    return normalized


def _ocr_pdf_page(file_path: str, page_number: int) -> str:
    try:
        from pdf2image import convert_from_path
    except ImportError as exc:
        raise DocumentExtractionError(
            "This PDF contains scanned pages. Install the OCR packages with "
            "`pip install pdf2image pytesseract pillow`, install the Tesseract "
            "OCR application and Poppler for Windows, then configure "
            "`TESSERACT_CMD` and `POPPLER_PATH` (or add both executables to PATH)."
        ) from exc

    try:
        import pytesseract
    except ImportError as exc:
        raise DocumentExtractionError(
            "This PDF contains scanned pages. Install the OCR packages with "
            "`pip install pdf2image pytesseract pillow`, install the Tesseract "
            "OCR application and Poppler for Windows, then configure "
            "`TESSERACT_CMD` and `POPPLER_PATH` (or add both executables to PATH)."
        ) from exc

    tesseract_cmd = os.environ.get("TESSERACT_CMD")
    if tesseract_cmd:
        if not os.path.isfile(tesseract_cmd):
            raise DocumentExtractionError(
                f"TESSERACT_CMD points to a missing executable: {tesseract_cmd}"
            )
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    elif not shutil.which("tesseract"):
        raise DocumentExtractionError(
            "Scanned PDF pages require the Tesseract OCR executable. Install "
            "Tesseract for Windows and add tesseract.exe to PATH, or set "
            "`TESSERACT_CMD` to its full path."
        )

    poppler_path = os.environ.get("POPPLER_PATH") or None
    if poppler_path and not os.path.isdir(poppler_path):
        raise DocumentExtractionError(
            f"POPPLER_PATH is not a directory: {poppler_path}"
        )
    if not poppler_path and not (
        shutil.which("pdftoppm") or shutil.which("pdftocairo")
    ):
        raise DocumentExtractionError(
            "Scanned PDF pages require Poppler. Install Poppler for Windows "
            "and add its `Library\\bin` directory to PATH, or set "
            "`POPPLER_PATH` to that directory."
        )

    images = []
    try:
        images = convert_from_path(
            file_path,
            dpi=OCR_DPI,
            first_page=page_number,
            last_page=page_number,
            fmt="png",
            grayscale=True,
            size=1600,
            thread_count=1,
            poppler_path=poppler_path,
            timeout=PDF_RENDER_TIMEOUT_SECONDS,
        )
        if not images:
            return ""
        return pytesseract.image_to_string(
            images[0],
            timeout=OCR_PAGE_TIMEOUT_SECONDS,
        ).strip()
    except DocumentExtractionError:
        raise
    except Exception as exc:
        error_name = type(exc).__name__
        if "PDFInfoNotInstalled" in error_name or "Poppler" in str(exc):
            raise DocumentExtractionError(
                "Poppler could not render this scanned PDF. Install Poppler "
                "for Windows and configure `POPPLER_PATH` or PATH."
            ) from exc
        if "TesseractNotFound" in error_name:
            raise DocumentExtractionError(
                "Tesseract was not found. Install Tesseract for Windows and "
                "configure `TESSERACT_CMD` or PATH."
            ) from exc
        if "Timeout" in error_name or "timeout" in str(exc).lower():
            raise DocumentExtractionError(
                f"OCR timed out on PDF page {page_number}. Try a lower-resolution "
                "or shorter document; OCR is limited to 30 seconds per page."
            ) from exc
        raise DocumentExtractionError(
            f"OCR failed on PDF page {page_number} ({error_name})."
        ) from exc
    finally:
        for image in images:
            image.close()


def extract_pdf_document(file_path: str) -> Dict[str, Any]:
    file_path = _validate_file_path(file_path)
    try:
        reader = PdfReader(file_path, strict=False)
        if reader.is_encrypted:
            try:
                decrypted = reader.decrypt("")
            except Exception:
                decrypted = 0
            if not decrypted:
                raise DocumentExtractionError(
                    "This PDF is encrypted or password-protected. Remove its "
                    "password protection before uploading."
                )
        page_count = len(reader.pages)
        if page_count == 0:
            raise DocumentExtractionError("The PDF contains no pages.")
        if page_count > MAX_PDF_PAGES:
            raise DocumentExtractionError(
                f"This PDF has {page_count} pages; the current safe limit is "
                f"{MAX_PDF_PAGES} pages."
            )
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(
            "The PDF is corrupt or could not be opened."
        ) from exc

    pages: List[Tuple[int, str]] = []
    ocr_pages: List[int] = []
    extracted_length = 0
    try:
        for page_number, page in enumerate(reader.pages, start=1):
            try:
                text = (page.extract_text() or "").strip()
            except Exception as exc:
                raise DocumentExtractionError(
                    f"Could not extract text from PDF page {page_number}."
                ) from exc
            if len(text) < MIN_PAGE_TEXT_LENGTH:
                embedded_text = text
                ocr_pages.append(page_number)
                ocr_text = _ocr_pdf_page(file_path, page_number)
                if ocr_text:
                    text = "\n".join(part for part in (embedded_text, ocr_text) if part)
            extracted_length += len(text)
            if extracted_length > MAX_TEXT_LENGTH:
                raise DocumentExtractionError(
                    "The extracted document text exceeds the 2,000,000-character "
                    "processing limit."
                )
            pages.append((page_number, text))
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(
            "The PDF could not be read. It may be corrupt or unsupported."
        ) from exc

    joined = "\n".join(
        f"--- Page {page_number} ---\n{text}"
        for page_number, text in pages
        if text
    )
    if not joined.strip():
        raise DocumentExtractionError(
            "No readable text was found in this PDF after text extraction and "
            "OCR. It may be blank, contain poor-quality scans, or require a "
            "different OCR language pack."
        )
    if len(joined) > MAX_TEXT_LENGTH:
        raise DocumentExtractionError(
            "The extracted document text exceeds the 2,000,000-character "
            "processing limit."
        )

    text_pages = [number for number, text in pages if text]
    empty_pages = [number for number, text in pages if not text]
    return {
        "text": joined,
        "summary": {
            "format": "pdf",
            "page_count": page_count,
            "pages_with_text": text_pages,
            "pages_ocr_processed": ocr_pages,
            "pages_without_text": empty_pages,
            "method": "text_and_ocr" if ocr_pages else "text",
            "warnings": (
                [f"No text could be extracted from PDF page(s): {', '.join(map(str, empty_pages))}."]
                if empty_pages
                else []
            ),
        },
    }


def extract_pdf_text(file_path: str) -> str:
    return extract_pdf_document(file_path)["text"]


def extract_docx_document(file_path: str) -> Dict[str, Any]:
    file_path = _validate_file_path(file_path)
    try:
        document = Document(file_path)
    except Exception as exc:
        raise DocumentExtractionError(
            "The DOCX is corrupt or could not be opened."
        ) from exc

    paragraphs = []
    total_length = 0
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            total_length += len(paragraph.text)
            if total_length > MAX_TEXT_LENGTH:
                raise DocumentExtractionError(
                    "The extracted document text exceeds the 2,000,000-character "
                    "processing limit."
                )
            paragraphs.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                row_text = " | ".join(cells)
                total_length += len(row_text)
                if total_length > MAX_TEXT_LENGTH:
                    raise DocumentExtractionError(
                        "The extracted document text exceeds the 2,000,000-character "
                        "processing limit."
                    )
                paragraphs.append(row_text)
    text = "\n".join(paragraphs).strip()
    if not text:
        raise DocumentExtractionError("No readable text was found in this DOCX.")
    if len(text) > MAX_TEXT_LENGTH:
        raise DocumentExtractionError(
            "The extracted document text exceeds the 2,000,000-character "
            "processing limit."
        )
    return {
        "text": text,
        "summary": {
            "format": "docx",
            "page_count": None,
            "pages_with_text": [],
            "pages_ocr_processed": [],
            "pages_without_text": [],
            "method": "text",
            "warnings": [],
        },
    }


def extract_docx_text(file_path: str) -> str:
    return extract_docx_document(file_path)["text"]


def extract_document(file_path: str) -> Dict[str, Any]:
    if not isinstance(file_path, str) or not file_path.strip():
        raise DocumentExtractionError("No file path was provided.")
    lower_path = file_path.lower()
    if lower_path.endswith(".pdf"):
        return extract_pdf_document(file_path)
    if lower_path.endswith(".docx"):
        return extract_docx_document(file_path)
    raise DocumentExtractionError("Unsupported file format. Upload a PDF or DOCX.")


def extract_text(file_path: str) -> str:
    return extract_document(file_path)["text"]
