import io
import os
import sys
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
from PIL import Image
from docx import Document
from pypdf import PdfReader, PdfWriter

from backend.auth import create_admin, initialize_database
from backend.main import app
from backend.services.document_parser import (
    DocumentExtractionError,
    extract_document,
    extract_text,
)


ROOT_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SAMPLE_PDF = os.path.join(ROOT_DIR, "data", "sample_contracts", "sample_contract.pdf")


class OCRIngestionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.auth_temp_dir = tempfile.TemporaryDirectory()
        cls.previous_database = os.environ.get("AUTH_DATABASE_PATH")
        os.environ["AUTH_DATABASE_PATH"] = os.path.join(
            cls.auth_temp_dir.name, "accounts.sqlite3"
        )
        os.environ["AUTH_COOKIE_SECURE"] = "false"
        initialize_database()
        create_admin("ocrtestadmin", "ocr@example.test", "ocr-test-password-long")

    @classmethod
    def tearDownClass(cls):
        if cls.previous_database is None:
            os.environ.pop("AUTH_DATABASE_PATH", None)
        else:
            os.environ["AUTH_DATABASE_PATH"] = cls.previous_database
        os.environ.pop("AUTH_COOKIE_SECURE", None)
        cls.auth_temp_dir.cleanup()

    def setUp(self):
        self.temp_paths = []
        self.client = TestClient(app, follow_redirects=False)
        csrf = self.client.get("/api/auth/csrf").json()["csrf_token"]
        login = self.client.post(
            "/api/auth/login",
            headers={"X-CSRF-Token": csrf},
            json={
                "username": "ocrtestadmin",
                "password": "ocr-test-password-long",
            },
        )
        self.assertEqual(login.status_code, 200, login.text)
        self.csrf = login.json()["csrf_token"]

    def tearDown(self):
        self.client.close()
        for path in self.temp_paths:
            if os.path.exists(path):
                os.remove(path)

    def _temp_path(self, suffix):
        handle = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
        handle.close()
        self.temp_paths.append(handle.name)
        return handle.name

    def _image_pdf(self):
        path = self._temp_path(".pdf")
        image = Image.new("RGB", (1200, 1600), "white")
        image.save(path, "PDF", resolution=150)
        image.close()
        return path

    def _mixed_pdf(self):
        path = self._temp_path(".pdf")
        image_pdf = self._image_pdf()
        writer = PdfWriter()
        writer.append_pages_from_reader(PdfReader(SAMPLE_PDF))
        writer.append_pages_from_reader(PdfReader(image_pdf))
        with open(path, "wb") as output:
            writer.write(output)
        return path

    def test_regular_text_pdf_uses_embedded_text(self):
        result = extract_document(SAMPLE_PDF)
        self.assertIn("SERVICE AGREEMENT", result["text"])
        self.assertEqual(result["summary"]["method"], "text")
        self.assertEqual(result["summary"]["pages_ocr_processed"], [])

    def test_pdf_page_limit_accepts_60_pages_and_rejects_61(self):
        path = self._temp_path(".pdf")

        class TextPage:
            def extract_text(self):
                return "This page contains enough selectable text to skip OCR."

        class Reader:
            is_encrypted = False

            def __init__(self, page_count):
                self.pages = [TextPage() for _ in range(page_count)]

        with patch(
            "backend.services.document_parser.PdfReader",
            side_effect=[Reader(60), Reader(61)],
        ):
            accepted = extract_document(path)
            self.assertEqual(accepted["summary"]["page_count"], 60)

            with self.assertRaisesRegex(DocumentExtractionError, "current safe limit is 60 pages"):
                extract_document(path)

    def test_image_only_pdf_uses_ocr_and_retains_page_marker(self):
        path = self._image_pdf()
        with patch(
            "backend.services.document_parser._ocr_pdf_page",
            return_value="This scanned agreement contains payment terms and termination.",
        ) as ocr:
            result = extract_document(path)

        self.assertIn("--- Page 1 ---", result["text"])
        self.assertIn("scanned agreement", result["text"])
        self.assertEqual(result["summary"]["method"], "text_and_ocr")
        self.assertEqual(result["summary"]["pages_ocr_processed"], [1])
        ocr.assert_called_once_with(path, 1)

    def test_short_embedded_text_is_preserved_when_ocr_returns_nothing(self):
        path = self._image_pdf()

        class ShortTextPage:
            def extract_text(self):
                return "Page 1"

        class Reader:
            is_encrypted = False
            pages = [ShortTextPage()]

        with patch("backend.services.document_parser.PdfReader", return_value=Reader()):
            with patch(
                "backend.services.document_parser._ocr_pdf_page",
                return_value="",
            ):
                result = extract_document(path)

        self.assertIn("Page 1", result["text"])
        self.assertEqual(result["summary"]["pages_ocr_processed"], [1])

    def test_mixed_pdf_ocr_runs_only_on_pages_without_text(self):
        path = self._mixed_pdf()
        with patch(
            "backend.services.document_parser._ocr_pdf_page",
            return_value="OCR page contains an arbitration provision and governing law.",
        ) as ocr:
            result = extract_document(path)

        self.assertIn("SERVICE AGREEMENT", result["text"])
        self.assertIn("OCR page contains an arbitration provision", result["text"])
        self.assertEqual(result["summary"]["pages_ocr_processed"], [2])
        ocr.assert_called_once_with(path, 2)

    def test_docx_regression_includes_paragraphs_and_tables(self):
        path = self._temp_path(".docx")
        document = Document()
        document.add_paragraph("Service Agreement")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Payment"
        table.cell(0, 1).text = "Invoices payable within 30 days"
        document.save(path)

        text = extract_text(path)
        self.assertIn("Service Agreement", text)
        self.assertIn("Payment | Invoices payable within 30 days", text)

    def test_empty_and_corrupt_pdf_fail_clearly(self):
        empty_path = self._temp_path(".pdf")
        with self.assertRaisesRegex(DocumentExtractionError, "corrupt"):
            extract_text(empty_path)

        corrupt_path = self._temp_path(".pdf")
        with open(corrupt_path, "wb") as output:
            output.write(b"%PDF-not-a-valid-document")
        with self.assertRaisesRegex(DocumentExtractionError, "corrupt"):
            extract_text(corrupt_path)

    def test_scanned_pdf_reports_missing_ocr_dependency(self):
        path = self._image_pdf()
        with patch.dict(sys.modules, {"pdf2image": None}):
            with self.assertRaisesRegex(DocumentExtractionError, "Install the OCR packages"):
                extract_text(path)

    def test_scanned_pdf_reports_missing_tesseract_executable(self):
        path = self._image_pdf()
        with patch.dict(os.environ, {"TESSERACT_CMD": ""}):
            with patch("backend.services.document_parser.shutil.which", return_value=None):
                with self.assertRaisesRegex(
                    DocumentExtractionError,
                    "require the Tesseract OCR executable",
                ):
                    extract_text(path)

    def test_ocr_timeout_is_reported(self):
        path = self._image_pdf()
        with patch(
            "backend.services.document_parser._ocr_pdf_page",
            side_effect=DocumentExtractionError(
                "OCR timed out on PDF page 1. Try a lower-resolution or shorter document."
            ),
        ):
            with self.assertRaisesRegex(DocumentExtractionError, "OCR timed out"):
                extract_text(path)

    def test_api_returns_extraction_summary_and_analysis(self):
        with open(SAMPLE_PDF, "rb") as source:
            response = self.client.post(
                "/api/analyze",
                headers={"X-CSRF-Token": self.csrf},
                files={
                    "file": (
                        "sample_contract.pdf",
                        source.read(),
                        "application/pdf",
                    )
                },
            )
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()

        self.assertGreater(result["text_length"], 0)
        self.assertGreater(result["extraction_summary"]["page_count"], 0)
        self.assertTrue(result["clauses"])
        self.assertTrue(result["compliance"])
        self.assertIn("overall_risk", result["risk"])
        self.assertIn("items", result["financial_impact"])
        self.assertIn("legal_analysis", result)
        self.assertEqual(result["legal_analysis"]["analysis_status"], "complete_with_review_required")
        self.assertIn("overview", result["legal_analysis"])
        self.assertIn("findings", result["legal_analysis"])
        self.assertIn("legal_sources", result["legal_analysis"])

    def test_api_sends_ocr_text_through_analysis_pipeline(self):
        path = self._image_pdf()
        with open(path, "rb") as source:
            scanned_pdf = source.read()
        recognized_text = (
            "SERVICE AGREEMENT. Payment terms require invoices within 30 days. "
            "Termination permits either party to end this agreement on notice."
        )
        with patch(
            "backend.services.document_parser._ocr_pdf_page",
            return_value=recognized_text,
        ):
            response = self.client.post(
                "/api/analyze",
                headers={"X-CSRF-Token": self.csrf},
                files={"file": ("scanned_contract.pdf", scanned_pdf, "application/pdf")},
            )
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()

        self.assertEqual(result["extraction_summary"]["pages_ocr_processed"], [1])
        self.assertGreater(result["text_length"], 0)
        self.assertTrue(result["clauses"])
        self.assertTrue(result["compliance"])
        self.assertIn("overall_risk", result["risk"])
        self.assertIn("items", result["financial_impact"])
        self.assertIn("legal_analysis", result)
        self.assertEqual(result["legal_analysis"]["analysis_status"], "complete_with_review_required")
        self.assertIn("overview", result["legal_analysis"])
        self.assertIn("findings", result["legal_analysis"])

    def test_api_law_analysis_uses_the_uploaded_cross_border_docx(self):
        path = self._temp_path(".docx")
        contract_text = (
            "This Software Services Agreement is between Beispiel GmbH, a company "
            "incorporated in Germany (Provider), and Bharat Tech Private Limited, "
            "incorporated in India (Customer).\n"
            "This Agreement shall be governed by the laws of Germany.\n"
            "The courts of Berlin shall have exclusive jurisdiction.\n"
            "Provider processes personal data in India."
        )
        document = Document()
        for paragraph in contract_text.split("\n"):
            document.add_paragraph(paragraph)
        document.save(path)

        with open(path, "rb") as source:
            response = self.client.post(
                "/api/analyze",
                headers={"X-CSRF-Token": self.csrf},
                files={
                    "file": (
                        "cross_border_services.docx",
                        source.read(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )

        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        legal = result["legal_analysis"]
        self.assertEqual(legal["governing_law"], "Germany")
        self.assertEqual(
            legal["overview"]["governing_law"]["jurisdiction"], "Germany"
        )
        self.assertEqual(
            [
                (party["name"], party["role"], party["country"])
                for party in legal["overview"]["parties"]
            ],
            [
                ("Beispiel GmbH", "Provider", "Germany"),
                ("Bharat Tech Private Limited", "Customer", "India"),
            ],
        )
        self.assertTrue(legal["overview"]["dispute_resolution"]["court_jurisdiction_evidence"])
        data_findings = [
            finding
            for finding in legal["findings"]
            if finding["contract_topic"] == "Data protection and transfers"
            and "processes personal data in India" in finding["contract_clause_text"]
        ]
        self.assertEqual(len(data_findings), 1)
        self.assertEqual(data_findings[0]["affected_party"], "Beispiel GmbH")
        self.assertEqual(data_findings[0]["jurisdiction"], "India")
        self.assertEqual(legal["source_verification"]["verified_provisions"], [])
        self.assertIn("risk", result)
        self.assertIn("compliance", result)
        self.assertIn("negotiation", result)
        self.assertIn("financial_impact", result)

    def test_api_rejects_empty_upload(self):
        response = self.client.post(
            "/api/analyze",
            headers={"X-CSRF-Token": self.csrf},
            files={"file": ("empty.pdf", io.BytesIO(b""), "application/pdf")},
        )
        self.assertEqual(response.status_code, 400)


if __name__ == "__main__":
    unittest.main()
